"""
Report generation: Excel (xlsx via openpyxl) and PDF (via reportlab).

Both functions return raw bytes so the routers can stream the file as a
StreamingResponse / Response with the appropriate Content-Disposition.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from ..models.print_log import PrintLog


REPORT_HEADERS: List[str] = [
    "ID",
    "Received At",
    "Username",
    "Computer",
    "Laptop IP",
    "Printer",
    "Printer IP",
    "Document",
    "Pages",
    "Copies",
    "Total Pages",
    "Status",
]


def _row(log: PrintLog) -> List:
    return [
        log.id,
        log.received_at.strftime("%Y-%m-%d %H:%M:%S") if log.received_at else "",
        log.username or "",
        log.computer_name or "",
        log.laptop_ip or "",
        log.printer_name or "",
        log.printer_ip or "",
        log.document_name or "",
        log.pages or 0,
        log.copies or 0,
        log.total_pages or 0,
        (log.status.value if hasattr(log.status, "value") else str(log.status)),
    ]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def build_excel_report(logs: Iterable[PrintLog]) -> bytes:
    """Return an .xlsx file (as bytes) containing the supplied print logs."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Print Logs"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    center = Alignment(horizontal="center", vertical="center")

    ws.append(REPORT_HEADERS)
    for col_idx, _ in enumerate(REPORT_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for log in logs:
        ws.append(_row(log))

    # Reasonable column widths
    widths = [6, 22, 18, 18, 16, 28, 16, 40, 8, 8, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def build_pdf_report(
    logs: Iterable[PrintLog],
    *,
    title: str = "Print Logs Report",
    subtitle: str | None = None,
) -> bytes:
    """Return a PDF report (as bytes) listing the supplied print logs."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=title,
    )

    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph(f"<b>{title}</b>", styles["Title"]))
    if subtitle:
        story.append(Paragraph(subtitle, styles["Normal"]))
    story.append(
        Paragraph(
            f"Generated at: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC",
            styles["Normal"],
        )
    )
    story.append(Spacer(1, 6 * mm))

    data: List[List] = [REPORT_HEADERS]
    for log in logs:
        data.append(_row(log))

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 6),
            ]
        )
    )
    story.append(table)

    doc.build(story)
    buf.seek(0)
    return buf.read()
